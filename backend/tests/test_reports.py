"""Tests for the reports module.

Checks:
- context builder produces a non-empty scorecard with takeaways
- PPTX file is a valid pptx with exactly TOTAL_SLIDES slides
- DOCX file is a valid docx with the expected section headings
- API endpoint returns the right content-type and headers
"""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient

from app.db import SessionLocal
from app.main import app
from reports.context import build_single_fund_context
from reports.generator import generate_report
from reports.pptx_report import TOTAL_SLIDES


@pytest.fixture(scope="module")
def client() -> TestClient:
    return TestClient(app)


def _skip_if_empty(client: TestClient) -> None:
    stats = client.get("/api/stats").json()
    if stats.get("funds", 0) == 0:
        pytest.skip("database not seeded; run `make ingest` first")


SUBJECT_CODE = 122639  # Parag Parikh Flexi Cap


# ---------- Context ----------


def test_context_shape() -> None:
    with SessionLocal() as s:
        if s.query(__import__("app.models", fromlist=["Fund"]).Fund).count() == 0:
            pytest.skip("db not seeded")
        ctx = build_single_fund_context(s, SUBJECT_CODE, period="3Y")
    assert ctx.fund.scheme_code == SUBJECT_CODE
    assert ctx.scorecard.volatility.available
    assert len(ctx.nav_series) > 500          # ~3 years of trading days
    assert len(ctx.calendar_years) >= 10
    assert any(c.bench_return is not None for c in ctx.calendar_years)
    assert 0 < (ctx.peer_rank_3y or 0) <= 100
    assert len(ctx.takeaways) == 3


# ---------- PPTX ----------


def test_pptx_structure() -> None:
    with SessionLocal() as s:
        if s.query(__import__("app.models", fromlist=["Fund"]).Fund).count() == 0:
            pytest.skip("db not seeded")
        r = generate_report(s, codes=[SUBJECT_CODE], fmt="pptx", period="3Y")

    assert r.content_type.endswith("presentation")
    assert r.filename.endswith(".pptx")
    assert len(r.body) > 50_000  # charts embedded

    from pptx import Presentation
    prs = Presentation(io.BytesIO(r.body))
    assert len(prs.slides) == TOTAL_SLIDES
    # Spot-check: first slide carries the fund name somewhere
    first_text = " ".join(
        sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame
    )
    assert "Parag Parikh" in first_text


# ---------- DOCX ----------


EXPECTED_HEADINGS = {
    "Executive summary", "Trailing returns", "Calendar-year returns",
    "Risk metrics", "Drawdown over time", "Rolling 3Y CAGR",
    "Peer ranking", "Methodology and disclaimer",
}


def test_docx_structure() -> None:
    with SessionLocal() as s:
        if s.query(__import__("app.models", fromlist=["Fund"]).Fund).count() == 0:
            pytest.skip("db not seeded")
        r = generate_report(s, codes=[SUBJECT_CODE], fmt="docx", period="3Y")
    assert r.content_type.endswith("wordprocessingml.document")
    assert r.filename.endswith(".docx")

    import docx
    doc = docx.Document(io.BytesIO(r.body))
    headings = {
        p.text for p in doc.paragraphs if p.style.name.startswith("Heading")
    }
    missing = EXPECTED_HEADINGS - headings
    assert not missing, f"missing section headings: {missing}"
    assert len(doc.tables) >= 4  # trailing + risk-summary + risk + peer tables
    # At least one embedded image (charts)
    assert len(doc.inline_shapes) >= 1


# ---------- API ----------


def test_reports_outline_endpoint(client: TestClient) -> None:
    r = client.get("/api/reports/outline")
    assert r.status_code == 200
    slides = r.json()["slides"]
    assert len(slides) == TOTAL_SLIDES
    assert slides[0]["title"] == "Title"
    assert slides[-1]["title"] == "Methodology & disclaimer"


def test_reports_preview(client: TestClient) -> None:
    _skip_if_empty(client)
    r = client.get(f"/api/reports/preview?codes={SUBJECT_CODE}&period=3Y")
    assert r.status_code == 200
    body = r.json()
    assert body["fund"]["scheme_code"] == SUBJECT_CODE
    assert len(body["takeaways"]) == 3


def test_reports_pptx_download(client: TestClient) -> None:
    _skip_if_empty(client)
    r = client.post(
        "/api/reports",
        json={"codes": [SUBJECT_CODE], "format": "pptx", "period": "3Y"},
    )
    assert r.status_code == 200
    assert r.headers["content-type"].endswith("presentation")
    assert "attachment" in r.headers["content-disposition"]
    assert r.headers.get("x-report-filename", "").endswith(".pptx")
    assert len(r.content) > 50_000


def test_reports_docx_download(client: TestClient) -> None:
    _skip_if_empty(client)
    r = client.post(
        "/api/reports",
        json={"codes": [SUBJECT_CODE], "format": "docx", "period": "3Y"},
    )
    assert r.status_code == 200
    assert r.headers["content-type"].endswith("wordprocessingml.document")


def test_reports_404_on_unknown_fund(client: TestClient) -> None:
    r = client.post(
        "/api/reports",
        json={"codes": [9999999], "format": "pptx", "period": "3Y"},
    )
    assert r.status_code == 404
