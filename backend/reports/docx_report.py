"""DOCX report builder.

Same content as the PPTX deck, as a structured document — headings, a
simple TOC, styled tables and embedded charts.
"""
from __future__ import annotations

import io
from datetime import date

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from reports import charts
from reports.context import ReportContext
from reports.tokens import (
    ACCENT, FONT_SANS,
    NEG, NEUTRAL_400, NEUTRAL_600, NEUTRAL_900, POS,
)


def build_docx(ctx: ReportContext) -> bytes:
    doc = Document()
    _configure_base_styles(doc)
    _configure_page(doc)

    _title_block(doc, ctx)
    _toc(doc)
    _executive_summary(doc, ctx)
    _performance_section(doc, ctx)
    _trailing_section(doc, ctx)
    _calendar_section(doc, ctx)
    _risk_section(doc, ctx)
    _drawdown_section(doc, ctx)
    _rolling_section(doc, ctx)
    _peer_section(doc, ctx)
    _methodology_section(doc, ctx)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Base styles
# ---------------------------------------------------------------------------


def _rgb(c: tuple[int, int, int]) -> RGBColor:
    return RGBColor(*c)


def _configure_page(doc: Document) -> None:
    for s in doc.sections:
        s.left_margin = Inches(0.9)
        s.right_margin = Inches(0.9)
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)


def _configure_base_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_SANS
    normal.font.size = Pt(11)
    normal.font.color.rgb = _rgb(NEUTRAL_900)

    for h, size in (("Heading 1", 20), ("Heading 2", 14)):
        style = doc.styles[h]
        style.font.name = FONT_SANS
        style.font.size = Pt(size)
        style.font.color.rgb = _rgb(NEUTRAL_900)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14 if h == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(4)


def _heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = FONT_SANS
        run.font.color.rgb = _rgb(NEUTRAL_900)


def _para(doc: Document, text: str, *, size: int = 11, color: tuple[int, int, int] = NEUTRAL_900,
          bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_SANS
    run.font.size = Pt(size)
    run.font.color.rgb = _rgb(color)
    run.bold = bold
    run.italic = italic


def _label(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.font.name = FONT_SANS
    run.font.size = Pt(9)
    run.font.color.rgb = _rgb(NEUTRAL_400)
    run.bold = False
    p.paragraph_format.space_after = Pt(2)


def _hr(doc: Document, color: tuple[int, int, int] = ACCENT, width_pt: int = 2) -> None:
    """Horizontal accent underline — a very thin box in a borderless table."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(0.8)
    cell = tbl.cell(0, 0)
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "%02X%02X%02X" % color)
    tc_pr.append(shd)
    # Collapse height to a hairline
    tr_pr = cell._tc.getparent().get_or_add_trPr()
    tr_h = OxmlElement("w:trHeight")
    tr_h.set(qn("w:val"), str(width_pt * 12))
    tr_h.set(qn("w:hRule"), "exact")
    tr_pr.append(tr_h)
    cell.text = ""


def _tone_colour(v: float | None) -> tuple[int, int, int]:
    if v is None:
        return NEUTRAL_900
    if v > 0:
        return POS
    if v < 0:
        return NEG
    return NEUTRAL_900


def _add_chart_image(doc: Document, png_bytes: bytes, width_in: float = 6.5) -> None:
    stream = io.BytesIO(png_bytes)
    doc.add_picture(stream, width=Inches(width_in))


def _table_style(table) -> None:
    table.style = "Light List"
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.name = FONT_SANS
                    run.font.size = Pt(10)


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------


def _title_block(doc: Document, ctx: ReportContext) -> None:
    _label(doc, "Fund scorecard")
    p = doc.add_paragraph()
    r = p.add_run(ctx.display_name)
    r.font.name = FONT_SANS
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = _rgb(NEUTRAL_900)

    sub = f"{ctx.fund.amc or '—'} · benchmarked against {ctx.fund.benchmark_id or 'none'} · {ctx.period} window"
    _para(doc, sub, size=11, color=NEUTRAL_600)
    _para(doc, f"Prepared {date.today():%d %b %Y}   ·   AlphaDesk", size=9, color=NEUTRAL_400)
    _hr(doc)


def _toc(doc: Document) -> None:
    _heading(doc, "Contents", level=2)
    for i, t in enumerate([
        "Executive summary",
        "Performance vs benchmark",
        "Trailing returns",
        "Calendar-year returns",
        "Risk metrics",
        "Drawdown over time",
        "Rolling 3Y CAGR",
        "Peer ranking",
        "Methodology and disclaimer",
    ], start=1):
        _para(doc, f"   {i}.  {t}", size=11, color=NEUTRAL_600)


def _executive_summary(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Executive summary")
    for t in ctx.takeaways:
        _para(doc, f"•  {t}")


def _performance_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, f"Performance vs {ctx.fund.benchmark_id or 'benchmark'}")
    _para(doc, f"Both series rebased to 100 at {ctx.nav_series.index[0].strftime('%d %b %Y')}.",
          color=NEUTRAL_600, size=10)
    png = charts.rebased_chart(
        ctx.nav_series,
        ctx.bench_series,
        fund_label=ctx.fund.name.split(" - ")[0],
        bench_label=ctx.fund.benchmark_id or "Benchmark",
        width_in=7.0,
        height_in=3.1,
    )
    _add_chart_image(doc, png)


def _trailing_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Trailing returns")
    headers = ["Metric", "1M", "3M", "6M", "1Y", "3Y", "5Y"]
    tr = ctx.scorecard.trailing

    table = doc.add_table(rows=2, cols=len(headers))
    _table_style(table)
    for i, h in enumerate(headers):
        c = table.cell(0, i)
        c.text = h
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = _rgb(NEUTRAL_600)

    table.cell(1, 0).text = ctx.display_name.split(" - ")[0]
    for i, k in enumerate(["1M", "3M", "6M", "1Y", "3Y", "5Y"], start=1):
        m = tr.get(k)
        text = "—" if not m or not m.available else f"{m.value * 100:+.2f}%"
        c = table.cell(1, i)
        c.text = text
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in p.runs:
                run.font.color.rgb = _rgb(_tone_colour(m.value if m and m.available else None))

    doc.add_paragraph()
    _label(doc, "Risk-adjusted summary")
    sc = ctx.scorecard
    rrows = [
        ("Volatility (ann.)", _fmt_pct(sc.volatility.value)),
        ("Max drawdown",       _fmt_pct(sc.max_drawdown.value, signed=True)),
        ("Sharpe",             _fmt_num(sc.sharpe.value)),
        ("Sortino",            _fmt_num(sc.sortino.value)),
    ]
    rtable = doc.add_table(rows=len(rrows), cols=2)
    _table_style(rtable)
    for i, (label, value) in enumerate(rrows):
        rtable.cell(i, 0).text = label
        c = rtable.cell(i, 1)
        c.text = value
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _calendar_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Calendar-year returns")
    if not ctx.calendar_years:
        _para(doc, "No calendar-year data available in the window.", color=NEUTRAL_600)
        return
    years_tuples = [(c.year, c.fund_return, c.bench_return) for c in ctx.calendar_years]
    png = charts.calendar_year_bars(years_tuples, width_in=7.0, height_in=3.0)
    _add_chart_image(doc, png)


def _risk_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Risk metrics")
    sc = ctx.scorecard
    rows = [
        ("Beta",              _fmt_num(sc.beta.value)),
        ("Jensen's α (ann.)", _fmt_pct(sc.alpha.value, signed=True)),
        ("Tracking error",    _fmt_pct(sc.tracking_error.value)),
        ("Information ratio", _fmt_num(sc.information_ratio.value)),
        ("Up capture",        _fmt_num(sc.up_capture.value)),
        ("Down capture",      _fmt_num(sc.down_capture.value)),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    _table_style(table)
    for i, (label, value) in enumerate(rows):
        table.cell(i, 0).text = label
        c = table.cell(i, 1)
        c.text = value
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _drawdown_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Drawdown over time")
    if ctx.drawdown_series.empty:
        _para(doc, "Not available.", color=NEUTRAL_600)
        return
    png = charts.drawdown_chart(ctx.drawdown_series, width_in=7.0, height_in=2.4)
    _add_chart_image(doc, png)
    if ctx.worst_drawdown is not None:
        wd = ctx.worst_drawdown
        rec = wd.recovery_date.strftime("%d %b %Y") if wd.recovery_date else "Not yet recovered"
        _para(doc,
              f"Worst drawdown: {wd.max_drawdown * 100:+.2f}%   ·   "
              f"Peak {wd.peak_date:%d %b %Y}   ·   Trough {wd.trough_date:%d %b %Y}   ·   Recovered {rec}",
              size=10, color=NEUTRAL_600)


def _rolling_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Rolling 3Y CAGR")
    if ctx.rolling_3y.empty:
        _para(doc, "Not available.", color=NEUTRAL_600)
        return
    png = charts.rolling_chart(
        ctx.rolling_3y,
        fund_label=ctx.fund.name.split(" - ")[0],
        bench_label=ctx.fund.benchmark_id or "Benchmark",
        width_in=7.0,
        height_in=3.0,
    )
    _add_chart_image(doc, png)


def _peer_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Peer ranking")
    _para(doc,
          f"Within {ctx.fund.category.replace('_', ' ').title()} ({len(ctx.peer_leaderboard)} funds).",
          color=NEUTRAL_600, size=10)

    table = doc.add_table(rows=len(ctx.peer_leaderboard) + 1, cols=3)
    _table_style(table)
    for i, h in enumerate(["Fund", "Code", "3Y CAGR"]):
        c = table.cell(0, i)
        c.text = h
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if i >= 1 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.bold = True
                run.font.color.rgb = _rgb(NEUTRAL_600)
    for i, p in enumerate(ctx.peer_leaderboard, start=1):
        marker = "  ◆" if p.is_subject else ""
        table.cell(i, 0).text = p.name.split(" - ")[0] + marker
        table.cell(i, 1).text = str(p.scheme_code)
        c = table.cell(i, 2)
        c.text = f"{p.return_3y * 100:.2f}%"
        for para in c.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for run in para.runs:
                run.font.color.rgb = _rgb(_tone_colour(p.return_3y))


def _methodology_section(doc: Document, ctx: ReportContext) -> None:
    _heading(doc, "Methodology and disclaimer")
    _para(doc, "Methodology")
    for b in [
        "Returns are computed from daily NAVs sourced from mfapi.in.",
        "Benchmark data: Nifty price indices via yfinance.",
        "Trailing returns < 1Y are absolute; ≥ 1Y are CAGR using calendar-day gaps.",
        "Risk statistics use sample standard deviation (ddof=1), 252 trading days per year.",
        "Risk-free rate for Sharpe / Sortino: 6.50% p.a. unless otherwise configured.",
        "Peer ranking is recomputed every ingest against the fund's category peers.",
    ]:
        _para(doc, f"•  {b}", size=10, color=NEUTRAL_600)

    _para(doc, "Disclaimer", bold=True)
    for d in ctx.disclaimers:
        _para(doc, f"•  {d}", size=10, color=NEUTRAL_600)


# ---------------------------------------------------------------------------
# formatters
# ---------------------------------------------------------------------------


def _fmt_pct(v: float | None, signed: bool = False) -> str:
    if v is None:
        return "—"
    return f"{v * 100:+.2f}%" if signed else f"{v * 100:.2f}%"


def _fmt_num(v: float | None) -> str:
    if v is None:
        return "—"
    return f"{v:.2f}"
